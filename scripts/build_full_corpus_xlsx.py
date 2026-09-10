"""
Snapshot the full category-only corpus (data/{group}_full_category_corpus.jsonl)
into a single Excel workbook, safe to re-run at any point while the underlying
pull (build_full_category_corpus.py) is still in progress -- it just reflects
whatever's on disk right now, clearly labeled with a live count and timestamp
rather than presented as a finished total.

Usage:
    python3 scripts/build_full_corpus_xlsx.py
"""
import json
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUT_PATH = DATA_DIR / "full_category_corpus_snapshot.xlsx"

GROUPS = ["quantum", "cs"]
COLUMNS = ["arxiv_id", "title", "authors", "published", "year",
           "primary_category", "abs_link", "pdf_link"]


def load_group(group: str) -> list[dict]:
    path = DATA_DIR / f"{group}_full_category_corpus.jsonl"
    records = []
    if not path.exists():
        return records
    with open(path) as f:
        for line in f:
            line = line.strip()
            if line:
                records.append(json.loads(line))
    return records


def main():
    counts = {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for group in GROUPS:
        records = load_group(group)
        counts[group] = len(records)
        ws = wb.create_sheet(title=group)
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in records:
            ws.append([
                r.get("arxiv_id", ""),
                r.get("title", ""),
                "; ".join(r.get("authors", [])),
                r.get("published", ""),
                r.get("year", ""),
                r.get("primary_category", ""),
                r.get("abs_link", ""),
                r.get("pdf_link", ""),
            ])
        ws.freeze_panes = "A2"
        ws.column_dimensions["B"].width = 70
        ws.column_dimensions["C"].width = 40

    done_flags = {}
    for group in GROUPS:
        cp_path = ROOT / "cache" / f"full_category_corpus_{group}_checkpoint.json"
        done_flags[group] = json.loads(cp_path.read_text()).get("done", False) if cp_path.exists() else False
    all_done = all(done_flags.values())

    summary = wb.create_sheet(title="Summary", index=0)
    title = "Full category-only corpus (quant-ph/cond-mat + CS categories, 2021-present, no keyword filter)"
    summary.append([title if all_done else title + " -- IN PROGRESS, not yet complete"])
    summary["A1"].font = Font(bold=True, size=13)
    summary.append([f"Snapshot taken: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
    summary.append([])
    summary.append(["Side", "Papers", "Status"])
    for cell in summary[4]:
        cell.font = Font(bold=True)
    for group in GROUPS:
        status = "Complete" if done_flags[group] else "In progress"
        summary.append([group, counts[group], status])
    summary.append([])
    summary.append(["Total (both sides)", sum(counts.values()), "Complete" if all_done else "In progress"])
    summary.column_dimensions["A"].width = 20
    summary.column_dimensions["C"].width = 16

    wb.save(OUT_PATH)
    print(f"wrote snapshot to {OUT_PATH}")
    for group in GROUPS:
        print(f"  {group}: {counts[group]} papers so far")


if __name__ == "__main__":
    main()
