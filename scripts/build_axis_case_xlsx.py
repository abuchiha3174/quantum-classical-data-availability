"""
Render the {group}_axis_case_review.csv produced by build_axis_case_csv.py as
a color-coded .xlsx, so a human reviewer can see at a glance which rows need
attention without opening every "Flag For Human" cell.

Rows with Flag For Human == "Yes" are filled red. Everything else (including
plain "pending_claude_code_classification" rows not yet reached by Stage 4)
is left with the default background.

This is a read-only rendering step -- it never edits the CSV, which remains
the single source of truth. Re-run this any time after build_axis_case_csv.py
to refresh the .xlsx.

Usage:
    python3 scripts/build_axis_case_xlsx.py quantum
    python3 scripts/build_axis_case_xlsx.py cs
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from domain_definitions import GROUPS  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"

FLAG_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER_FONT = Font(bold=True)

# Columns wide enough to skim; long free-text columns are capped so the sheet
# stays navigable, with wrap_text on so nothing is silently truncated.
COLUMN_WIDTHS = {
    "Title": 40, "Authors": 25, "Published Date": 14, "arXiv ID": 12,
    "Paper Link": 30, "Axis A Case": 22, "Axis A Evidence": 50,
    "Axis B Case": 22, "Axis B Evidence": 50, "Confidence": 12,
    "Flag For Human": 14, "Notes": 40, "Link(s) Checked": 30,
    "Link Verification Result": 30, "Verified": 12,
    "Corrected Axis A Case": 20, "Corrected Axis B Case": 20, "Reviewer Notes": 30,
}


def main():
    if len(sys.argv) != 2 or sys.argv[1] not in GROUPS:
        print(f"usage: python3 {Path(__file__).name} [{'|'.join(GROUPS)}]")
        sys.exit(1)
    group = sys.argv[1]

    csv_path = DATA_DIR / f"{group}_axis_case_review.csv"
    if not csv_path.exists():
        print(f"{csv_path} not found -- run build_axis_case_csv.py {group} first")
        sys.exit(1)

    with open(csv_path, newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        print(f"{csv_path} is empty")
        sys.exit(1)
    header, body = rows[0], rows[1:]
    flag_col = header.index("Flag For Human")

    wb = Workbook()
    ws = wb.active
    ws.title = f"{group}_axis_case_review"

    ws.append(header)
    for col_idx, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(name, 18)
    ws.freeze_panes = "A2"

    n_flagged = 0
    for row in body:
        ws.append(row)
        r = ws.max_row
        if row[flag_col].strip().lower() == "yes":
            n_flagged += 1
            for col_idx in range(1, len(header) + 1):
                ws.cell(row=r, column=col_idx).fill = FLAG_FILL
        for col_idx in range(1, len(header) + 1):
            ws.cell(row=r, column=col_idx).alignment = Alignment(vertical="top", wrap_text=True)

    out_path = DATA_DIR / f"{group}_axis_case_review.xlsx"
    wb.save(out_path)
    print(f"wrote {len(body)} rows to {out_path}")
    print(f"{n_flagged}/{len(body)} rows highlighted red (Flag For Human = Yes)")


if __name__ == "__main__":
    main()
